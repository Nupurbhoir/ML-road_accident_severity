import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_report_docx():
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    report_path = 'Project_Report.md'
    if not os.path.exists(report_path):
        print("Project_Report.md not found!")
        return

    with open(report_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_text = []
    in_table = False
    table_data = []

    def flush_table():
        nonlocal in_table, table_data
        if not table_data:
            return
        
        cols_count = max(len(row) for row in table_data)
        table = doc.add_table(rows=len(table_data), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_value in enumerate(row):
                if c_idx < cols_count:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_value.strip()
                    
                    # Formatting Header Row
                    if r_idx == 0:
                        set_cell_background(cell, '1E293B')
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, 'F8FAFC')
                        else:
                            set_cell_background(cell, 'FFFFFF')
                            
                    cell.paragraphs[0].paragraph_format.space_before = Pt(4)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(4)
                    
        doc.add_paragraph()
        table_data = []
        in_table = False

    def flush_code():
        nonlocal in_code_block, code_text
        if not code_text:
            return
        code_str = "".join(code_text)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(code_str)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        code_text = []
        in_code_block = False

    for line in lines:
        stripped = line.strip()
        
        # Check for Code Blocks
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            continue
            
        if in_code_block:
            code_text.append(line)
            continue
            
        # Check for Markdown Tables
        if '|' in stripped and ('---' in stripped or not stripped.startswith('!')):
            if not in_table:
                in_table = True
                table_data = []
            if '---' in stripped:
                continue
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            if cells:
                table_data.append(cells)
            continue
        else:
            if in_table:
                flush_table()

        if not stripped:
            continue

        # Check for Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[2:])
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:])
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:])
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        # Check for Images
        elif stripped.startswith('!['):
            img_match = re.search(r'!\[.*?\]\((.*?)\)', stripped)
            if img_match:
                img_path = img_match.group(1)
                if os.path.exists(img_path):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(8)
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(5.8))
                else:
                    print(f"Warning: Image {img_path} not found.")
        # Bullet Points
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            text = stripped[2:]
            # Simple bold parser **text**
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    output_filename = 'Project_Report.docx'
    doc.save(output_filename)
    print(f"Successfully generated Word Document: {output_filename}")

if __name__ == "__main__":
    create_report_docx()
